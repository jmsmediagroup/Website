#!/usr/bin/env python3
"""Write the free ADHD budget template to assets/downloads/adhd-budget-planner/ADHD-Budget-Template.xlsx.

    python3 make_budget_xlsx.py

Follows The ADHD Budget Planner's own structure, so the spreadsheet and the
book give the same answer: a Bills Master List with due days and an autopay
column, a one-page Money Map, a ten-minute weekly check-in, and an ADHD Tax
log. Amounts carry no currency symbol, because readers use several.

Needs openpyxl. Every total is a formula, so the file works in Excel, Google
Sheets (File, Import, Upload) and Numbers without anything precomputed.
openpyxl writes formulas without cached results, so recalculate once (open and
save in LibreOffice or Excel) before committing; otherwise phone previews,
which do not calculate, show the totals as blank.
"""

import os

from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.datavalidation import DataValidation

OUT = os.path.join(os.path.dirname(os.path.abspath(__file__)), "assets", "downloads", "adhd-budget-planner",
                   "ADHD-Budget-Template.xlsx")

FONT = "Arial"
INPUT = Font(name=FONT, size=11, color="0000FF")     # numbers you change
CALC = Font(name=FONT, size=11, color="000000")      # formulas
LINK = Font(name=FONT, size=11, color="008000")      # pulled from another sheet
BODY = Font(name=FONT, size=11)
BOLD = Font(name=FONT, size=11, bold=True)
TITLE = Font(name=FONT, size=16, bold=True)
NOTE = Font(name=FONT, size=10, italic=True, color="555555")
HEAD_FILL = PatternFill("solid", start_color="F1EDE4")
KEY_FILL = PatternFill("solid", start_color="FFF6CC")
LINE = Border(bottom=Side(style="thin", color="CCCCCC"))
MONEY = '#,##0.00;(#,##0.00);"-"'
WRAP = Alignment(wrap_text=True, vertical="top")

BILL_ROWS = 25
WEEKS = 52
LOG_ROWS = 60


def header(ws, row, labels, widths):
    for col, (label, width) in enumerate(zip(labels, widths), start=1):
        c = ws.cell(row=row, column=col, value=label)
        c.font, c.fill, c.border = BOLD, HEAD_FILL, LINE
        ws.column_dimensions[c.column_letter].width = width


def start_sheet(wb):
    ws = wb.active
    ws.title = "Start here"
    ws.column_dimensions["A"].width = 92
    lines = [
        ("The ADHD Budget Template", TITLE),
        ("Free from jmstruijk.com. Follows the structure of The ADHD Budget Planner.", NOTE),
        ("", BODY),
        ("One hour, once", BOLD),
        ("1. Bills: list every bill with its amount, the day it is due, and whether it is on autopay.", BODY),
        ("2. Money Map: enter what comes in each month and a figure for each flexible category.", BODY),
        ("3. Move anything marked No onto autopay, if you can. Each one is a due date you no longer have to remember.", BODY),
        ("", BODY),
        ("Ten minutes, once a week", BOLD),
        ("4. Weekly check-in: write what you spent this week. The sheet compares it with your weekly money.", BODY),
        ("5. ADHD Tax log: add any late fee, forgotten subscription or replacement buy when it happens.", BODY),
        ("", BODY),
        ("If you miss a week, or a month", BOLD),
        ("Nothing breaks. The weeks are numbered rather than dated, so start again at the next row.", BODY),
        ("", BODY),
        ("Blue numbers are yours to change. Black numbers are worked out for you. Green ones come from another sheet.", NOTE),
        ("Amounts are in whatever currency you use. Nothing in this file is sent anywhere.", NOTE),
    ]
    for i, (text, font) in enumerate(lines, start=1):
        c = ws.cell(row=i, column=1, value=text)
        c.font, c.alignment = font, WRAP


def bills_sheet(wb):
    ws = wb.create_sheet("Bills")
    ws["A1"], ws["A1"].font = "Bills master list", TITLE
    ws["A2"], ws["A2"].font = "Every bill, with the day it is due. Getting these out of your head is the point.", NOTE
    header(ws, 4, ["Bill", "Per month", "Due day (1-31)", "On autopay?", "Notes"], [30, 14, 16, 14, 40])
    first, last = 5, 4 + BILL_ROWS
    for r in range(first, last + 1):
        ws.cell(row=r, column=2).number_format = MONEY
        for col in (2, 3, 4):
            ws.cell(row=r, column=col).font = INPUT
        ws.cell(row=r, column=1).font = BODY
        ws.cell(row=r, column=5).font = BODY

    yes_no = DataValidation(type="list", formula1='"Yes,No"', allow_blank=True)
    day = DataValidation(type="whole", operator="between", formula1="1", formula2="31", allow_blank=True,
                         error="Enter a day of the month from 1 to 31.", showErrorMessage=True)
    ws.add_data_validation(yes_no)
    ws.add_data_validation(day)
    yes_no.add("D%d:D%d" % (first, last))
    day.add("C%d:C%d" % (first, last))

    t = last + 2
    ws.cell(row=t, column=1, value="Total bills per month").font = BOLD
    total = ws.cell(row=t, column=2, value="=SUM(B%d:B%d)" % (first, last))
    total.font, total.number_format, total.fill = Font(name=FONT, size=11, bold=True), MONEY, KEY_FILL
    ws.cell(row=t + 1, column=1, value="Bills not on autopay").font = BOLD
    manual = ws.cell(row=t + 1, column=2,
                     value='=COUNTIFS(A%d:A%d,"<>",D%d:D%d,"<>Yes")' % (first, last, first, last))
    manual.font = CALC
    manual.comment = Comment("Counts named bills not marked Yes. Each one is a date you have to remember.", "jmstruijk.com")
    ws.freeze_panes = "A5"
    return "Bills!$B$%d" % t


def money_map_sheet(wb, bills_total):
    ws = wb.create_sheet("Money Map")
    ws.column_dimensions["A"].width = 38
    ws.column_dimensions["B"].width = 16
    ws.column_dimensions["C"].width = 58
    ws["A1"], ws["A1"].font = "Money map", TITLE
    ws["A2"], ws["A2"].font = "One page for the month. Few categories on purpose: every extra one is another decision.", NOTE
    header(ws, 4, ["", "Per month", "Note"], [38, 16, 58])

    rows = [
        (5, "Take-home income", None, INPUT, "What actually arrives, after tax."),
        (7, "Bills", "=" + bills_total, LINK, "From the Bills sheet."),
        (8, "Groceries", None, INPUT, ""),
        (9, "Getting around", None, INPUT, "Fuel, fares, parking."),
        (10, "Everything else", None, INPUT, "Kept as one figure. Splitting it further rarely survives a month."),
        (11, "ADHD tax buffer", None, INPUT,
         "Late fees, forgotten subscriptions, replacements. Plan for it. "
         "The ADHD Tax Calculator at jmstruijk.com/tools/ gives your own figure."),
        (12, "Savings", None, INPUT, "Moved automatically on payday, ideally."),
    ]
    for r, label, formula, font, note in rows:
        ws.cell(row=r, column=1, value=label).font = BODY
        c = ws.cell(row=r, column=2, value=formula)
        c.font, c.number_format = font, MONEY
        ws.cell(row=r, column=3, value=note).font = NOTE
        ws.cell(row=r, column=3).alignment = WRAP

    ws.cell(row=14, column=1, value="Total going out").font = BOLD
    out = ws.cell(row=14, column=2, value="=SUM(B7:B12)")
    out.font, out.number_format = Font(name=FONT, size=11, bold=True), MONEY

    ws.cell(row=15, column=1, value="Left over").font = BOLD
    left = ws.cell(row=15, column=2, value="=B5-B14")
    left.font, left.number_format, left.fill = Font(name=FONT, size=11, bold=True), MONEY, KEY_FILL
    ws.cell(row=15, column=3, value='=IF(B5=0,"Enter your income first.",IF(B15<0,"The plan spends more than '
                                    'comes in. Trim a flexible figure.","Balanced."))').font = NOTE

    ws.cell(row=17, column=1, value="Weeks in a month").font = BODY
    weeks = ws.cell(row=17, column=2, value="=52/12")
    weeks.font, weeks.number_format = CALC, "0.00"
    ws.cell(row=17, column=3, value="52 weeks over 12 months, so a month is a little over four weeks.").font = NOTE

    ws.cell(row=18, column=1, value="Weekly spending money").font = BOLD
    weekly = ws.cell(row=18, column=2, value="=SUM(B8:B10)/B17")
    weekly.font, weekly.number_format, weekly.fill = Font(name=FONT, size=11, bold=True), MONEY, KEY_FILL
    ws.cell(row=18, column=3, value="Groceries, getting around and everything else, per week. "
                                    "A week is easier to keep track of than a month.").font = NOTE
    ws.cell(row=18, column=3).alignment = WRAP
    return "'Money Map'!$B$18"


def weekly_sheet(wb, weekly_money):
    ws = wb.create_sheet("Weekly check-in")
    ws["A1"], ws["A1"].font = "Weekly check-in", TITLE
    ws["A2"], ws["A2"].font = "Ten minutes. Numbered rather than dated, so a missed week costs nothing.", NOTE
    header(ws, 4, ["Week", "Week starting", "Spent this week", "Weekly money", "Difference", "One thing to change"],
           [8, 16, 16, 16, 14, 46])
    for i in range(1, WEEKS + 1):
        r = 4 + i
        ws.cell(row=r, column=1, value=i).font = BODY
        ws.cell(row=r, column=2).font = INPUT
        spent = ws.cell(row=r, column=3)
        spent.font, spent.number_format = INPUT, MONEY
        budget = ws.cell(row=r, column=4, value="=" + weekly_money)
        budget.font, budget.number_format = LINK, MONEY
        diff = ws.cell(row=r, column=5, value='=IF(C%d="","",D%d-C%d)' % (r, r, r))
        diff.font, diff.number_format = CALC, MONEY
        ws.cell(row=r, column=6).font = BODY
    ws.freeze_panes = "A5"


def tax_log_sheet(wb):
    ws = wb.create_sheet("ADHD Tax log")
    ws["A1"], ws["A1"].font = "ADHD Tax log", TITLE
    ws["A2"], ws["A2"].font = "Write it down when it happens, without judging it. The pattern is what is useful.", NOTE
    header(ws, 4, ["Date", "What happened", "Type", "Amount", "Could autopay or a reminder stop it?"],
           [14, 40, 24, 14, 36])
    first, last = 5, 4 + LOG_ROWS
    kinds = DataValidation(type="list", allow_blank=True,
                           formula1='"Late fee,Forgotten subscription,Impulse buy,Replacement,Convenience,Other"')
    ws.add_data_validation(kinds)
    kinds.add("C%d:C%d" % (first, last))
    for r in range(first, last + 1):
        for col in (1, 2, 3, 5):
            ws.cell(row=r, column=col).font = BODY
        amt = ws.cell(row=r, column=4)
        amt.font, amt.number_format = INPUT, MONEY

    t = last + 2
    ws.cell(row=t, column=1, value="Total so far").font = BOLD
    total = ws.cell(row=t, column=4, value="=SUM(D%d:D%d)" % (first, last))
    total.font, total.number_format, total.fill = Font(name=FONT, size=11, bold=True), MONEY, KEY_FILL
    for i, kind in enumerate(["Late fee", "Forgotten subscription", "Impulse buy", "Replacement", "Convenience", "Other"]):
        r = t + 2 + i
        ws.cell(row=r, column=3, value=kind).font = BODY
        c = ws.cell(row=r, column=4, value='=SUMIF(C%d:C%d,"%s",D%d:D%d)' % (first, last, kind, first, last))
        c.font, c.number_format = CALC, MONEY
    ws.freeze_panes = "A5"


def main():
    wb = Workbook()
    start_sheet(wb)
    bills_total = bills_sheet(wb)
    weekly_money = money_map_sheet(wb, bills_total)
    weekly_sheet(wb, weekly_money)
    tax_log_sheet(wb)
    wb.active = 0
    os.makedirs(os.path.dirname(OUT), exist_ok=True)
    wb.save(OUT)
    print("wrote", OUT)


if __name__ == "__main__":
    main()
